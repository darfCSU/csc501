#!/usr/bin/env python3
"""Generate an APA 7 student-paper reference template for Pandoc."""

import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_font(style, size=12, bold=None, italic=None):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)

    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic

    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts

    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)

    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attribute}"), "Times New Roman")


def set_double_spacing(style):
    fmt = style.paragraph_format
    fmt.line_spacing = 2
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)


def set_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.first_line_indent = Inches(0)

    run = paragraph.add_run()
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")

    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")

    number = OxmlElement("w:t")
    number.text = "1"

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(separate)
    run._r.append(number)
    run._r.append(end)


def set_list_indentation(doc):
    """Prevent the Normal style's first-line indent from shifting lists."""
    for style_name in ("Compact", "List Paragraph"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            set_font(style)
            set_double_spacing(style)
            style.paragraph_format.left_indent = Inches(0)
            style.paragraph_format.first_line_indent = Inches(0)


def create_template(output_path):
    # Start with Pandoc's reference document to preserve its numbering.
    with tempfile.TemporaryDirectory() as temp_dir:
        base_template = Path(temp_dir) / "reference.docx"

        with base_template.open("wb") as template_file:
            subprocess.run(
                ["pandoc", "--print-default-data-file=reference.docx"],
                stdout=template_file,
                check=True,
            )

        doc = Document(base_template)

        # APA 7 student papers use one-inch margins.
        section = doc.sections[0]
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

        # Body text.
        normal = doc.styles["Normal"]
        set_font(normal)
        set_double_spacing(normal)
        normal.paragraph_format.first_line_indent = Inches(0.5)

        # Keep list paragraphs from inheriting the body-text first-line indent.
        set_list_indentation(doc)

        # APA Level 1: centered, bold.
        heading1 = doc.styles["Heading 1"]
        set_font(heading1, bold=True)
        set_double_spacing(heading1)
        heading1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading1.paragraph_format.first_line_indent = Inches(0)
        heading1.paragraph_format.keep_with_next = True

        # APA Level 2: flush left, bold.
        heading2 = doc.styles["Heading 2"]
        set_font(heading2, bold=True)
        set_double_spacing(heading2)
        heading2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        heading2.paragraph_format.first_line_indent = Inches(0)
        heading2.paragraph_format.keep_with_next = True

        # APA Level 3: flush left, bold italic.
        heading3 = doc.styles["Heading 3"]
        set_font(heading3, bold=True, italic=True)
        set_double_spacing(heading3)
        heading3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        heading3.paragraph_format.first_line_indent = Inches(0)
        heading3.paragraph_format.keep_with_next = True

        # APA Level 4: indented, bold, ending with a period.
        heading4 = doc.styles["Heading 4"]
        set_font(heading4, bold=True)
        set_double_spacing(heading4)
        heading4.paragraph_format.left_indent = Inches(0.5)
        heading4.paragraph_format.first_line_indent = Inches(0)
        heading4.paragraph_format.keep_with_next = True

        # APA Level 5: indented, bold italic, ending with a period.
        heading5 = doc.styles["Heading 5"]
        set_font(heading5, bold=True, italic=True)
        set_double_spacing(heading5)
        heading5.paragraph_format.left_indent = Inches(0.5)
        heading5.paragraph_format.first_line_indent = Inches(0)
        heading5.paragraph_format.keep_with_next = True

        # Title style.
        title = doc.styles["Title"]
        set_font(title, bold=True)
        set_double_spacing(title)
        title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.first_line_indent = Inches(0)

        # References style with a hanging indent.
        styles = doc.styles
        if "References" not in styles:
            references = styles.add_style(
                "References",
                1,  # WD_STYLE_TYPE.PARAGRAPH
            )
        else:
            references = styles["References"]

        set_font(references)
        set_double_spacing(references)
        references.paragraph_format.left_indent = Inches(0.5)
        references.paragraph_format.first_line_indent = Inches(-0.5)

        # Page number in the top-right header.
        header = section.header
        header.paragraphs[0].clear()
        set_page_number(header.paragraphs[0])

        # Save the reference document.
        output_path = Path(output_path).expanduser().resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)

    print(f"Created APA reference template: {output_path}")


def main():
    output = (
        Path(sys.argv[1])
        if len(sys.argv) == 2
        else Path(__file__).resolve().parent / "apa-reference.docx"
    )

    if len(sys.argv) > 2:
        print(
            f"Usage: {Path(sys.argv[0]).name} [output.docx]",
            file=sys.stderr,
        )
        return 2

    try:
        create_template(output)
    except (OSError, subprocess.CalledProcessError) as exc:
        print(f"Error creating template: {exc}", file=sys.stderr)
        return 1

    return 0


if __name__ == "__main__":
    sys.exit(main())
